import os
from typing import Dict, Any, List
import pandas as pd
import openpyxl
from sqlalchemy.orm import Session
from sqlalchemy import select, delete

from app.models import (
    Subsidiary,
    Mine,
    User,
    UserScope,
    ProductionAnnual,
    ProductionMonthly,
    DispatchSummary,
    CoalQuality,
    GeologicalUnit,
    MiningIssueLog,
    InspectionRegister,
    NationalAnnualProduction,
    CaptiveCommercialProduction,
    MacroStatisticalTable,
    ParliamentaryBenchmark
)
from app.ingestion.cleaner import load_clean_csv, clean_numeric_str, clean_int, clean_text


def seed_master_and_users(db: Session, coal_data_dir: str) -> Dict[str, int]:
    """Reads mine_master.csv and seeds subsidiaries, mines, demo users, and user scopes."""
    master_path = os.path.join(coal_data_dir, "mine_master.csv")
    df = load_clean_csv(master_path)
    
    sub_count = 0
    mine_count = 0

    # 1. Subsidiaries & Mines
    for _, row in df.iterrows():
        sub_name = clean_text(row.get("subsidiary"))
        sub_id = sub_name.upper().replace(" ", "_").replace(".", "")[:40]
        
        sub = db.execute(select(Subsidiary).where(Subsidiary.subsidiary_name == sub_name)).scalar_one_or_none()
        if not sub:
            sub = Subsidiary(subsidiary_id=sub_id, subsidiary_name=sub_name)
            db.add(sub)
            db.flush()
            sub_count += 1
            
        m_code = clean_text(row.get("mine_code"))
        mine = db.execute(select(Mine).where(Mine.mine_code == m_code)).scalar_one_or_none()
        if not mine:
            mine = Mine(
                mine_code=m_code,
                mine_name=clean_text(row.get("mine_name")),
                subsidiary_id=sub.subsidiary_id,
                mine_type=clean_text(row.get("mine_type")),
                coal_type=clean_text(row.get("coal_type")),
                location=clean_text(row.get("location"))
            )
            db.add(mine)
            mine_count += 1
        else:
            mine.mine_name = clean_text(row.get("mine_name"))
            mine.mine_type = clean_text(row.get("mine_type"))
            mine.coal_type = clean_text(row.get("coal_type"))
            mine.location = clean_text(row.get("location"))

    # 2. Canonical Demo Users (from AGENTS.md Section 32)
    demo_users = [
        {
            "user_id": "USR001",
            "username": "mining_eng_deom",
            "email": "usr001@cmpdi.co.in",
            "role": "Mining Engineer",
            "department": "Mining",
            "clearance_level": "INTERNAL",
            "assigned_mine_code": "DEOM-01",
            "scopes": [{"mine_code": "DEOM-01", "department": "Mining", "max_classification": "INTERNAL"}]
        },
        {
            "user_id": "USR002",
            "username": "geology_eng_deom",
            "email": "usr002@cmpdi.co.in",
            "role": "Geology Engineer",
            "department": "Geology",
            "clearance_level": "RESTRICTED",
            "assigned_mine_code": "DEOM-01",
            "scopes": [{"mine_code": "DEOM-01", "department": "Geology", "max_classification": "RESTRICTED"}]
        },
        {
            "user_id": "USR003",
            "username": "transport_eng_knug",
            "email": "usr003@cmpdi.co.in",
            "role": "Transportation Engineer",
            "department": "Transportation",
            "clearance_level": "INTERNAL",
            "assigned_mine_code": "KNUG-02",
            "scopes": [{"mine_code": "KNUG-02", "department": "Transportation", "max_classification": "INTERNAL"}]
        },
        {
            "user_id": "USR004",
            "username": "mine_manager_multi",
            "email": "usr004@cmpdi.co.in",
            "role": "Mine Manager",
            "department": "Executive",
            "clearance_level": "RESTRICTED",
            "assigned_mine_code": None,
            "scopes": [
                {"mine_code": "DEOM-01", "department": "ALL", "max_classification": "RESTRICTED"},
                {"mine_code": "KNUG-02", "department": "ALL", "max_classification": "RESTRICTED"},
            ]
        },
        {
            "user_id": "USR005",
            "username": "admin_hq",
            "email": "usr005@cmpdi.co.in",
            "role": "Administrator",
            "department": "Administration",
            "clearance_level": "CONFIDENTIAL",
            "assigned_mine_code": None,
            "scopes": [{"mine_code": None, "department": "ALL", "max_classification": "CONFIDENTIAL"}]
        }
    ]

    user_count = 0
    for udata in demo_users:
        u = db.execute(select(User).where(User.user_id == udata["user_id"])).scalar_one_or_none()
        if not u:
            u = User(
                user_id=udata["user_id"],
                username=udata["username"],
                email=udata["email"],
                role=udata["role"],
                department=udata["department"],
                clearance_level=udata["clearance_level"],
                assigned_mine_code=udata["assigned_mine_code"],
                is_active=True
            )
            db.add(u)
            db.flush()
            for sdata in udata["scopes"]:
                scope = UserScope(
                    user_id=u.user_id,
                    mine_code=sdata["mine_code"],
                    department=sdata["department"],
                    max_classification=sdata["max_classification"]
                )
                db.add(scope)
            user_count += 1

    db.commit()
    return {"subsidiaries": sub_count, "mines": mine_count, "users": user_count}


def ingest_operational_tables(db: Session, coal_data_dir: str) -> Dict[str, int]:
    """Ingests all operational CSV records into normalized tables."""
    counts = {}

    # A. Production Annual & Targets (Reconciled)
    prod_path = os.path.join(coal_data_dir, "production", "production_2021_2025.csv")
    targets_path = os.path.join(coal_data_dir, "targets", "annual_targets_2021_2025.csv")
    
    prod_df = load_clean_csv(prod_path)
    targets_df = load_clean_csv(targets_path)
    
    # Merge on mine_code and year
    merged = pd.merge(prod_df, targets_df, on=["mine_code", "year"], suffixes=("_prod", "_target"))
    
    pa_count = 0
    for _, row in merged.iterrows():
        m_code = clean_text(row["mine_code"])
        yr = clean_int(row["year"])
        
        record = db.execute(
            select(ProductionAnnual).where(
                ProductionAnnual.mine_code == m_code,
                ProductionAnnual.year == yr
            )
        ).scalar_one_or_none()
        
        if not record:
            record = ProductionAnnual(mine_code=m_code, year=yr, actual_production_mt=0.0)
            db.add(record)
            pa_count += 1
            
        record.target_mt = clean_numeric_str(row.get("annual_target_mt"))
        record.actual_production_mt = clean_numeric_str(row.get("actual_production_mt")) or clean_numeric_str(row.get("actual_mt"))
        record.variance_mt = clean_numeric_str(row.get("variance_mt"))
        record.achievement_pct = clean_numeric_str(row.get("achievement_pct")) or clean_numeric_str(row.get("target_achievement_pct"))
        record.dispatch_mt = clean_numeric_str(row.get("dispatch_mt"))
        record.equipment_or_face_availability_pct = clean_numeric_str(row.get("equipment_or_face_availability_pct"))
    
    counts["production_annual"] = len(merged)

    # B. Monthly Production (180 records)
    monthly_path = os.path.join(coal_data_dir, "monthly", "monthly_production.csv")
    m_df = load_clean_csv(monthly_path)
    for _, row in m_df.iterrows():
        m_code = clean_text(row["mine_code"])
        yr = clean_int(row["year"])
        mo = clean_int(row["month"])
        
        rec = db.execute(
            select(ProductionMonthly).where(
                ProductionMonthly.mine_code == m_code,
                ProductionMonthly.year == yr,
                ProductionMonthly.month == mo
            )
        ).scalar_one_or_none()
        
        if not rec:
            rec = ProductionMonthly(
                mine_code=m_code,
                year=yr,
                month=mo,
                production_mt=clean_numeric_str(row["production_mt"]),
                monthly_target_mt=clean_numeric_str(row["monthly_target_mt"]),
                variance_mt=clean_numeric_str(row["variance_mt"])
            )
            db.add(rec)
        else:
            rec.production_mt = clean_numeric_str(row["production_mt"])
            rec.monthly_target_mt = clean_numeric_str(row["monthly_target_mt"])
            rec.variance_mt = clean_numeric_str(row["variance_mt"])
    counts["production_monthly"] = len(m_df)

    # C. Dispatch Summary (15 records)
    disp_path = os.path.join(coal_data_dir, "dispatch", "dispatch_summary.csv")
    d_df = load_clean_csv(disp_path)
    for _, row in d_df.iterrows():
        m_code = clean_text(row["mine_code"])
        yr = clean_int(row["year"])
        
        rec = db.execute(
            select(DispatchSummary).where(
                DispatchSummary.mine_code == m_code,
                DispatchSummary.year == yr
            )
        ).scalar_one_or_none()
        
        if not rec:
            rec = DispatchSummary(
                mine_code=m_code,
                year=yr,
                dispatch_mt=clean_numeric_str(row["dispatch_mt"]),
                production_dispatch_gap_mt=clean_numeric_str(row["production_dispatch_gap_mt"]),
                mode=clean_text(row["mode"]),
                logistics_status=clean_text(row["logistics_status"])
            )
            db.add(rec)
        else:
            rec.dispatch_mt = clean_numeric_str(row["dispatch_mt"])
            rec.production_dispatch_gap_mt = clean_numeric_str(row["production_dispatch_gap_mt"])
            rec.mode = clean_text(row["mode"])
            rec.logistics_status = clean_text(row["logistics_status"])
    counts["dispatch_summary"] = len(d_df)

    # D. Coal Quality Summary (15 records)
    q_path = os.path.join(coal_data_dir, "quality", "coal_quality_summary.csv")
    q_df = load_clean_csv(q_path)
    for _, row in q_df.iterrows():
        m_code = clean_text(row["mine_code"])
        yr = clean_int(row["year"])
        
        rec = db.execute(
            select(CoalQuality).where(
                CoalQuality.mine_code == m_code,
                CoalQuality.year == yr
            )
        ).scalar_one_or_none()
        
        if not rec:
            rec = CoalQuality(
                mine_code=m_code,
                year=yr,
                ash_pct=clean_numeric_str(row["ash_pct"]),
                moisture_pct=clean_numeric_str(row["moisture_pct"]),
                quality_action=clean_text(row["quality_action"])
            )
            db.add(rec)
        else:
            rec.ash_pct = clean_numeric_str(row["ash_pct"])
            rec.moisture_pct = clean_numeric_str(row["moisture_pct"])
            rec.quality_action = clean_text(row["quality_action"])
    counts["coal_quality"] = len(q_df)

    # E. Geological Units (9 records)
    geo_path = os.path.join(coal_data_dir, "geology", "geological_units.csv")
    g_df = load_clean_csv(geo_path)
    for _, row in g_df.iterrows():
        m_code = clean_text(row["mine_code"])
        unit_str = clean_text(row["unit"])
        
        rec = db.execute(
            select(GeologicalUnit).where(
                GeologicalUnit.mine_code == m_code,
                GeologicalUnit.unit == unit_str
            )
        ).scalar_one_or_none()
        
        if not rec:
            rec = GeologicalUnit(
                mine_code=m_code,
                unit=unit_str,
                depth_or_horizon_m=clean_text(row["depth_or_horizon_m"]),
                thickness_m=clean_text(row["thickness_m"]),
                structure=clean_text(row["structure"]),
                geological_risk=clean_text(row["geological_risk"]),
                observation=clean_text(row["observation"])
            )
            db.add(rec)
        else:
            rec.depth_or_horizon_m = clean_text(row["depth_or_horizon_m"])
            rec.thickness_m = clean_text(row["thickness_m"])
            rec.structure = clean_text(row["structure"])
            rec.geological_risk = clean_text(row["geological_risk"])
            rec.observation = clean_text(row["observation"])
    counts["geological_units"] = len(g_df)

    # F. Mining Issue Log (12 records)
    iss_path = os.path.join(coal_data_dir, "issues", "mining_issue_log.csv")
    i_df = load_clean_csv(iss_path)
    # Clear and re-populate for clean idempotency
    db.execute(delete(MiningIssueLog))
    for _, row in i_df.iterrows():
        rec = MiningIssueLog(
            mine_code=clean_text(row["mine_code"]),
            year=clean_int(row["year"]),
            issue_category=clean_text(row["issue_category"]),
            observed_issue=clean_text(row["observed_issue"]),
            operational_impact=clean_text(row["operational_impact"]),
            corrective_action=clean_text(row["corrective_action"])
        )
        db.add(rec)
    counts["mining_issue_log"] = len(i_df)

    # G. Inspection Register (12 records)
    insp_path = os.path.join(coal_data_dir, "inspections", "inspection_register.csv")
    ins_df = load_clean_csv(insp_path)
    db.execute(delete(InspectionRegister))
    for _, row in ins_df.iterrows():
        rec = InspectionRegister(
            mine_code=clean_text(row["mine_code"]),
            year=clean_int(row["year"]),
            inspection_focus=clean_text(row["inspection_focus"]),
            status=clean_text(row["status"]),
            observation=clean_text(row["observation"]),
            responsible_officer=clean_text(row["responsible_officer"])
        )
        db.add(rec)
    counts["inspection_register"] = len(ins_df)

    db.commit()
    return counts


def ingest_macro_and_benchmarks(db: Session, coal_data_dir: str) -> Dict[str, int]:
    """Ingests official macro statistical CSVs, Excel workbooks, and golden benchmark Q&As."""
    counts = {}

    # 1. National Annual Production (11 records)
    nat_path = os.path.join(coal_data_dir, "Official Data", "Annual_Coal_Production.csv")
    if os.path.exists(nat_path):
        nat_df = load_clean_csv(nat_path)
        for _, row in nat_df.iterrows():
            yr_str = clean_text(row.get("Year") or row.get("﻿Year"))
            rec = db.execute(select(NationalAnnualProduction).where(NationalAnnualProduction.year_range == yr_str)).scalar_one_or_none()
            if not rec:
                rec = NationalAnnualProduction(
                    year_range=yr_str,
                    coking_coal_mt=clean_numeric_str(row["Coking_Coal_MT"]),
                    non_coking_coal_mt=clean_numeric_str(row["Non_Coking_Coal_MT"]),
                    growth_pct=clean_numeric_str(row["Growth"]),
                    total_mt=clean_numeric_str(row["Total_MT"])
                )
                db.add(rec)
            else:
                rec.coking_coal_mt = clean_numeric_str(row["Coking_Coal_MT"])
                rec.non_coking_coal_mt = clean_numeric_str(row["Non_Coking_Coal_MT"])
                rec.growth_pct = clean_numeric_str(row["Growth"])
                rec.total_mt = clean_numeric_str(row["Total_MT"])
        counts["national_annual_production"] = len(nat_df)

    # 2. Captive & Commercial Mines (39 records)
    cap_path = os.path.join(coal_data_dir, "Official Data", "Coal-Production-of-Captive-and-Commercial-Coal-Mines.csv")
    if os.path.exists(cap_path):
        cap_df = load_clean_csv(cap_path)
        db.execute(delete(CaptiveCommercialProduction))
        for _, row in cap_df.iterrows():
            rec = CaptiveCommercialProduction(
                end_use=clean_text(row.get("End_Use") or row.get("﻿End_Use")),
                state=clean_text(row["State"]),
                company=clean_text(row["Company"]),
                quantity_mt=clean_numeric_str(row["Quantity_MT"])
            )
            db.add(rec)
        counts["captive_commercial_production"] = len(cap_df)

    # 3. Macro Statistical Tables (cdchap1.xlsx through cdchap5.xlsx)
    macro_count = 0
    official_dir = os.path.join(coal_data_dir, "Official Data")
    for fname in sorted(os.listdir(official_dir)):
        if fname.startswith("cdchap") and fname.endswith(".xlsx"):
            fpath = os.path.join(official_dir, fname)
            # extract chapter number from filename (e.g. cdchap1 -> 1)
            ch_num = int("".join([c for c in fname if c.isdigit()]) or 1)
            wb = openpyxl.load_workbook(fpath, read_only=True)
            for sname in wb.sheetnames:
                sheet = wb[sname]
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    # stringify row
                    row_vals = [str(c).strip() if c is not None else "" for c in row]
                    if any(row_vals):
                        rows.append(row_vals)
                
                title = rows[0][0] if rows and rows[0] and rows[0][0] else f"Table {sname}"
                
                rec = db.execute(
                    select(MacroStatisticalTable).where(
                        MacroStatisticalTable.source_file == fname,
                        MacroStatisticalTable.table_id == sname
                    )
                ).scalar_one_or_none()
                
                payload = {"rows": rows[:500]} # store rows in structured jsonb
                if not rec:
                    rec = MacroStatisticalTable(
                        source_file=fname,
                        chapter=ch_num,
                        table_id=sname,
                        table_title=title[:500],
                        data_json=payload
                    )
                    db.add(rec)
                    macro_count += 1
                else:
                    rec.table_title = title[:500]
                    rec.data_json = payload
    counts["macro_statistical_tables"] = macro_count

    # 4. Parliamentary Benchmarks (5 golden questions)
    pq_path = os.path.join(coal_data_dir, "parliamentary", "parliamentary_qa.csv")
    if os.path.exists(pq_path):
        pq_df = load_clean_csv(pq_path)
        for _, row in pq_df.iterrows():
            qid = clean_text(row["question_id"])
            rec = db.execute(select(ParliamentaryBenchmark).where(ParliamentaryBenchmark.question_id == qid)).scalar_one_or_none()
            if not rec:
                rec = ParliamentaryBenchmark(
                    question_id=qid,
                    question=clean_text(row["question"]),
                    reference_answer=clean_text(row["reference_answer"])
                )
                db.add(rec)
            else:
                rec.question = clean_text(row["question"])
                rec.reference_answer = clean_text(row["reference_answer"])
        counts["parliamentary_benchmarks"] = len(pq_df)

    db.commit()
    return counts
